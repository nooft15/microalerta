"""Pruebas de la API REST (Flask test client) con base de datos temporal."""
import io
import os
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from microalerta import create_app
from microalerta.seed import seed_db


@pytest.fixture()
def app(tmp_path):
    a = create_app({
        "TESTING": True,
        "DATABASE": str(tmp_path / "test.db"),
        "UPLOAD_DIR": str(tmp_path / "uploads"),
        "EXCEL_PATH": str(tmp_path / "datos.xlsx"),
        "SECRET_KEY": "test-secret",
    })
    seed_db(a)
    yield a


@pytest.fixture()
def client(app):
    return app.test_client()


def login(client, username="admin", password="demo123"):
    r = client.post("/api/login", json={"username": username, "password": password})
    assert r.status_code == 200, r.get_json()
    return r


def stub_weather(monkeypatch):
    """Evita llamadas de red: respuestas simuladas de Open-Meteo."""
    from microalerta import weather as w

    class R:
        def __init__(self, payload):
            self.payload = payload
        def raise_for_status(self):
            pass
        def json(self):
            return self.payload

    def fake_get(url, *args, **kwargs):
        if "archive" in url:
            return R({"hourly": {
                "temperature_2m": [20.0, 21.0],
                "relative_humidity_2m": [40.0, 42.0],
                "wind_speed_10m": [15.0, 17.0],
                "wind_direction_10m": [200, 210],
                "wind_gusts_10m": [30.0, 33.0],
                "pressure_msl": [1013.0, 1014.0],
                "precipitation": [0.0, 0.4]}})
        return R({})

    monkeypatch.setattr(w.requests, "get", fake_get)


def make_report(client, stain="lugol", n=12, files=None):
    payload = {
        "location_id": "cop-liceo",
        "dish_id": "PL-TEST",
        "sample_type": "muestra",
        "replicate_group": "RG-TEST",
        "exposure_start": "2026-01-05T09:00:00",
        "exposure_end": "2026-01-06T09:00:00",
        "volume_ml": "40",
        "area_cm2": "4.0",
        "medium": "Agua destilada congelada",
        "stain": stain,
        "microscope_mag": "400",
        "n_observed": str(n),
        "morph_fiber": "6",
        "morph_fragment": "4",
        "morph_film": "1",
        "morph_mineral": "1",
        "n_uncertain": "1",
        "obs_time": "10:30",
        "notes": "muestra de prueba",
    }
    # En esta versión de Werkzeug los archivos van dentro de `data`
    # como tuplas (fileobj, filename).
    for name, (fobj, fname) in (files or {}).items():
        payload[name] = (fobj, fname)
    return client.post("/api/reports", data=payload,
                       content_type="multipart/form-data")


def test_health_and_public_session(client):
    assert client.get("/api/health").get_json()["ok"] is True
    assert client.get("/api/session").get_json()["user"] is None


def test_login_roles_and_permissions(client):
    login(client, "admin")
    assert client.get("/api/session").get_json()["user"]["role"] == "admin"
    client.post("/api/logout")
    login(client, "ciudadano")
    assert client.get("/api/pending").status_code == 403
    assert client.get("/api/stats").status_code == 403
    assert client.get("/api/locations").status_code == 200


def test_index_served(client):
    r = client.get("/")
    assert r.status_code == 200
    assert "MicroAlerta" in r.get_data(as_text=True)


def test_stains_can_be_limited_by_admin(client):
    login(client, "admin")
    base = client.get("/api/stains?all=1").get_json()["stains"]
    assert {s["value"] for s in base} == {"lugol", "azul_metileno", "ninguna"}

    r = client.put("/api/stains", json={"allowed": ["lugol", "ninguna"]})
    assert r.status_code == 200
    assert {s["value"] for s in r.get_json()["stains"]} == {"lugol", "ninguna"}
    # El formulario público solo recibe las activas
    pub = client.get("/api/stains").get_json()["stains"]
    assert {s["value"] for s in pub} == {"lugol", "ninguna"}


def test_citizen_cannot_manage_stains(client):
    login(client, "ciudadano")
    assert client.put("/api/stains", json={"allowed": ["lugol"]}).status_code == 403


def test_report_requires_login(client):
    assert make_report(client).status_code == 401

def test_report_flow_and_admin_review(client, monkeypatch):
    stub_weather(monkeypatch)
    login(client, "colegio")
    r = make_report(client, n=20)
    assert r.status_code == 201, r.get_json()
    rep = r.get_json()["report"]
    assert rep["status"] == "pendiente"
    assert rep["metrics"]["conc_p_ml"] == pytest.approx(0.5, abs=1e-3)
    assert rep["semaphore"]["level"] >= 1
    obs_id = rep["obs_id"]

    # El clima de la ventana de exposición queda guardado en el detalle
    detail = client.get(f"/api/reports/{obs_id}").get_json()
    assert detail["weather"]["source"] == "open-meteo-archivo"
    assert detail["obs_time"] == "10:30"

    # Ciudadano no puede aprobar
    client.post("/api/logout")
    login(client, "ciudadano")
    assert client.post(f"/api/reports/{obs_id}/review",
                       json={"status": "aprobado"}).status_code == 403

    # Admin aprueba y aparece en comunidad
    client.post("/api/logout")
    login(client, "admin")
    r2 = client.post(f"/api/reports/{obs_id}/review",
                     json={"status": "aprobado", "comment": "Válido por conteo doble"})
    assert r2.status_code == 200
    assert r2.get_json()["report"]["status"] == "aprobado"

    comm = client.get("/api/community").get_json()
    assert any(l["report"]["obs_id"] == obs_id for l in comm["locations"])


def test_report_with_image_upload(client, monkeypatch):
    stub_weather(monkeypatch)
    login(client, "ciudadano")
    files = {"image": (io.BytesIO(b"fake-png-data"), "foto-placa.png")}
    r = make_report(client, stain="azul_metileno", n=5, files=files)
    assert r.status_code == 201, r.get_json()
    rep = r.get_json()["report"]
    assert rep["image_path"] and rep["image_path"].startswith("/static/uploads/")


def test_disallowed_stain_rejected(client):
    login(client, "colegio")
    r = make_report(client, stain="cristal_violeta")
    assert r.status_code == 400
    assert "tinción" in r.get_json()["error"]


def test_invalid_location_rejected(client):
    login(client, "colegio")
    payload = {
        "location_id": "no-existe", "sample_type": "muestra",
        "stain": "lugol", "n_observed": "3",
    }
    r = client.post("/api/reports", data=payload, content_type="multipart/form-data")
    assert r.status_code == 400


def test_selfcare_weather_endpoint(client):
    r = client.get("/api/selfcare", query_string={
        "conc": "6.5", "hum": "15", "wind": "50", "temp": "2"})
    assert r.status_code == 200
    text = " ".join(x["title"] + x["detail"] for x in r.get_json()["recommendations"])
    assert "N95" in text

    # Unidades preferidas: el viento del aviso se formatea en mph
    ru = client.get("/api/selfcare", query_string={
        "conc": "6.5", "hum": "15", "wind": "50", "temp": "32",
        "utemp": "F", "uwind": "mph"})
    assert ru.status_code == 200
    text_u = " ".join(x["detail"] for x in ru.get_json()["recommendations"])
    assert "mph" in text_u


def test_semaphore_scale_endpoint(client):
    r = client.get("/api/semaphore")
    assert r.status_code == 200
    levels = r.get_json()["levels"]
    assert len(levels) == 4
    for lv in levels:
        assert lv["color"].startswith("#")
        assert lv["recommendations"]


    rw = client.get("/api/weather")
    assert rw.status_code == 400  # faltan lat/lon


def test_community_aggregates_seeded(client):
    d = client.get("/api/community").get_json()
    assert d["total_approved"] >= 4
    assert d["cities"]["Copiapó"] >= 1
    assert d["counts_by_kind"].get("blanco", 0) >= 1
    first = d["locations"][0]["report"]
    assert "index" in first and "color" in first["index"]


def test_report_detail_with_qc(client):
    reports = client.get("/api/reports?status=aprobado").get_json()["reports"]
    g1 = next(r for r in reports
              if r["replicate_group"] == "RG-COP-01" and r["sample_type"] == "muestra")
    d = client.get(f"/api/reports/{g1['obs_id']}").get_json()
    assert d["qc_notes"] == [] or any("blanco" in x for x in d["qc_notes"])
    assert d["blank_qc"]["ok"] is True


# ── Funciones nuevas: CRUD admin, config video, prefs, stats ampliadas ──────
def test_precipitation_stored_with_snapshot(client, monkeypatch):
    stub_weather(monkeypatch)
    login(client, "colegio")
    r = make_report(client, n=10)
    obs_id = r.get_json()["report"]["obs_id"]
    d = client.get(f"/api/reports/{obs_id}").get_json()
    assert d["weather"]["precip_mm"] == pytest.approx(0.4)


def test_stats_extended_shape(client):
    login(client, "admin")
    d = client.get("/api/stats").get_json()
    assert "rejected" in d["stats"]
    assert isinstance(d["by_location"], list)
    assert set(d["by_morphology"]) == {"morph_fiber", "morph_fragment",
                                       "morph_film", "morph_mineral"}
    assert "lugol" in d["by_stain"]
    assert d["avg_conc_approved"] is not None


def test_video_config_default_and_admin_update(client):
    base = client.get("/api/config").get_json()["video_url"]
    assert base.startswith("http")

    login(client, "ciudadano")
    assert client.put("/api/config/video",
                      json={"url": "https://x.dev/v"}).status_code == 403
    client.post("/api/logout")

    login(client, "admin")
    r = client.put("/api/config/video", json={"url": "https://youtu.be/demo123"})
    assert r.status_code == 200
    assert client.get("/api/config").get_json()["video_url"] == "https://youtu.be/demo123"
    assert client.put("/api/config/video", json={"url": "no-es-url"}).status_code == 400


def test_prefs_roundtrip_and_auth(client):
    assert client.get("/api/prefs").status_code == 401
    assert client.put("/api/prefs", json={"units": {"temp": "F"}}).status_code == 401

    login(client, "colegio")
    r = client.put("/api/prefs", json={"units": {"temp": "F", "wind": "kn"}})
    assert r.status_code == 200
    units = r.get_json()["prefs"]["units"]
    assert units["temp"] == "F" and units["wind"] == "kn"
    assert units["conc"] == "pml"  # normalizado con el valor por defecto

    g = client.get("/api/prefs").get_json()["prefs"]["units"]
    assert g["temp"] == "F" and g["wind"] == "kn"


def test_excel_autosync_and_download(client, monkeypatch):
    path = client.application.config["EXCEL_PATH"]

    stub_weather(monkeypatch)
    login(client, "colegio")
    r = make_report(client, n=15)
    obs_id = r.get_json()["report"]["obs_id"]
    assert os.path.exists(path)          # el Excel se genera al registrar la muestra

    client.post("/api/logout")
    login(client, "admin")
    assert client.post(f"/api/reports/{obs_id}/review",
                       json={"status": "aprobado"}).status_code == 200
    assert os.path.exists(path)          # y se refresca al aprobar

    # Sin sesión no se puede descargar
    fresh = client.application.test_client()
    assert fresh.get("/api/export.xlsx").status_code == 401

    resp = client.get("/api/export.xlsx")
    assert resp.status_code == 200
    assert "spreadsheet" in resp.headers["Content-Type"]
    assert resp.data[:2] == b"PK"        # firma de archivo .xlsx (zip)

    import io as _io
    from openpyxl import load_workbook
    wb = load_workbook(_io.BytesIO(resp.data))
    assert {"Muestras", "Serie temporal", "Resumen por punto",
            "Morfologías", "Cómo usar"}.issubset(set(wb.sheetnames))
    ws = wb["Muestras"]
    assert ws.max_row >= 2               # encabezado + al menos una muestra
    headers = [c.value for c in ws[1]]
    assert "Concentración (p/mL)" in headers
    assert "Reporte" in headers


def test_admin_edit_and_delete_report(client, monkeypatch):
    stub_weather(monkeypatch)
    login(client, "colegio")
    r = make_report(client, n=12)
    obs_id = r.get_json()["report"]["obs_id"]

    # Ciudadano no puede editar ni eliminar
    client.post("/api/logout")
    login(client, "ciudadano")
    assert client.put(f"/api/reports/{obs_id}", json={"n_observed": 30}).status_code == 403
    assert client.delete(f"/api/reports/{obs_id}").status_code == 403

    # Admin edita: la concentración se recalcula
    client.post("/api/logout")
    login(client, "admin")
    ru = client.put(f"/api/reports/{obs_id}", json={"n_observed": 30})
    assert ru.status_code == 200
    assert ru.get_json()["report"]["metrics"]["conc_p_ml"] == pytest.approx(0.75)

    # Admin elimina: desaparece de listado y detalle
    rd = client.delete(f"/api/reports/{obs_id}")
    assert rd.status_code == 200 and rd.get_json()["ok"] is True
    assert client.get(f"/api/reports/{obs_id}").status_code == 404


def test_diagnose_endpoint(client):
    """Panel: el diagnóstico público enumera el estado del sistema."""
    r = client.get("/api/diagnose")
    assert r.status_code == 200
    d = r.get_json()
    assert d["ok"] is True
    names = [c["name"] for c in d["checks"]]
    assert "Base de datos" in names
    assert "Sesión" in names
    assert "Excel de datos" in names
    assert "Fuente de clima" in names       # en TESTING se omite sin red
    for c in d["checks"]:
        assert c["status"] in ("ok", "warn", "error")


def test_maintenance_requires_admin(client, monkeypatch):
    """Mantenimiento (backfill/excel) solo para admin."""
    fresh = client.application.test_client()
    assert fresh.post("/api/maintenance",
                      json={"action": "excel_sync"}).status_code == 401

    login(client, "colegio")
    assert client.post("/api/maintenance",
                       json={"action": "excel_sync"}).status_code == 403

    client.post("/api/logout")
    stub_weather(monkeypatch)
    login(client, "admin")
    rx = client.post("/api/maintenance", json={"action": "excel_sync"})
    assert rx.status_code == 200 and rx.get_json()["ok"] is True

    rb = client.post("/api/maintenance", json={"action": "backfill_weather"})
    assert rb.status_code == 200
    assert "backfill" in rb.get_json()

    # Acción desconocida -> 400
    assert client.post("/api/maintenance",
                       json={"action": "nada"}).status_code == 400